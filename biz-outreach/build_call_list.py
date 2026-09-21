"""Exports all no-website leads to a call-tracking spreadsheet."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import CONFIG
from storage import connect

OUTPUT_PATH = "exports/no_website_call_list.xlsx"

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
BODY_FONT = Font(name="Arial")
LEGEND_FONT = Font(name="Arial", italic=True, size=9, color="6B7280")

COLUMNS = [
    ("Business Name", 32),
    ("Category", 22),
    ("Phone", 16),
    ("Address", 40),
    ("Demo Link", 42),
    ("Called?", 10),
    ("Outcome", 20),
    ("Notes", 40),
]


def build():
    with connect(CONFIG.db_path) as conn:
        rows = conn.execute(
            "SELECT name, category, address, phone, demo_url FROM leads ORDER BY category, name"
        ).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Call List"

    ws.append(["No-website business call list - fill in Called?/Outcome/Notes as you work through it"])
    ws["A1"].font = LEGEND_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    header_row = 3
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for i, row in enumerate(rows):
        r = header_row + 1 + i
        category = (row["category"] or "").replace("_", " ").title()
        ws.cell(row=r, column=1, value=row["name"]).font = BODY_FONT
        ws.cell(row=r, column=2, value=category).font = BODY_FONT
        ws.cell(row=r, column=3, value=row["phone"] or "").font = BODY_FONT
        ws.cell(row=r, column=4, value=row["address"] or "").font = BODY_FONT
        ws.cell(row=r, column=5, value=row["demo_url"] or "(not hosted yet)").font = BODY_FONT
        ws.cell(row=r, column=6, value="").font = BODY_FONT
        ws.cell(row=r, column=7, value="").font = BODY_FONT
        ws.cell(row=r, column=8, value="").font = BODY_FONT

    last_row = header_row + len(rows)
    summary_row = last_row + 2
    ws.cell(row=summary_row, column=1, value="Total leads:").font = Font(name="Arial", bold=True)
    ws.cell(row=summary_row, column=2, value=f"=COUNTA(A{header_row + 1}:A{last_row})").font = BODY_FONT
    ws.cell(row=summary_row + 1, column=1, value="Called:").font = Font(name="Arial", bold=True)
    ws.cell(row=summary_row + 1, column=2, value=f'=COUNTIF(F{header_row + 1}:F{last_row},"Y")').font = BODY_FONT

    ws.freeze_panes = f"A{header_row + 1}"

    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} with {len(rows)} leads")


if __name__ == "__main__":
    build()
